# Templates for the two xtt PRs. pip install openpyxl xlsxwriter
# Each has A1 = "Title:" and B1 = "{R-TITLE}"; merge it with R-TITLE = 'ACME Corp'.
import openpyxl, xlsxwriter, zipfile

# PR 1: openpyxl writes Target="/xl/worksheets/sheet1.xml" in workbook.xml.rels
wb = openpyxl.Workbook()
ws = wb.active
ws["A1"] = "Title:"
ws["B1"] = "{R-TITLE}"
wb.save("openpyxl_template.xlsx")

# PR 2: XlsxWriter in constant_memory mode writes inline strings, no sharedStrings.xml
wb = xlsxwriter.Workbook("xlsxwriter_template.xlsx", {"constant_memory": True})
ws = wb.add_worksheet("Sheet1")
ws.write("A1", "Title:")
ws.write("B1", "{R-TITLE}")
wb.close()

# PR 1, isolated: a normal XlsxWriter file (shared strings, like Excel),
# only the sheet Target in workbook.xml.rels made absolute as openpyxl does
wb = xlsxwriter.Workbook("tmp.xlsx")
ws = wb.add_worksheet("Sheet1")
ws.write("A1", "Title:")
ws.write("B1", "{R-TITLE}")
wb.close()
with zipfile.ZipFile("tmp.xlsx") as src, zipfile.ZipFile("absolute_target_only.xlsx", "w", zipfile.ZIP_DEFLATED) as dst:
    for item in src.infolist():
        data = src.read(item.filename)
        if item.filename == "xl/_rels/workbook.xml.rels":
            data = data.replace(b'Target="worksheets/', b'Target="/xl/worksheets/')
        dst.writestr(item, data)
